# Topic 04: Styled Multi-Sheet Excel Automation
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_styled_excel(output_filename="sales_report.xlsx"):
    # Sample datasets
    df_summary = pd.DataFrame({
        "Department": ["Engineering", "Marketing", "Sales", "Support"],
        "Headcount": [45, 18, 30, 22],
        "Budget": [450000.0, 180000.0, 320000.0, 140000.0],
        "Utilization_%": [94.5, 88.0, 99.2, 91.0]
    })

    df_details = pd.DataFrame({
        "Employee": ["Alice W.", "Bob T.", "Carlos G.", "Diane V."],
        "Department": ["Engineering", "Marketing", "Sales", "Support"],
        "Project": ["Apollo", "Brand Refresh", "Enterprise Q3", "Helpdesk AI"]
    })

    with pd.ExcelWriter(output_filename, engine='openpyxl') as writer:
        df_summary.to_excel(writer, sheet_name='Summary', index=False)
        df_details.to_excel(writer, sheet_name='Team Details', index=False)

        wb = writer.book

        # Style Summary Sheet
        ws = wb['Summary']
        ws.freeze_panes = 'A2'

        navy_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

        for cell in ws[1]:
            cell.fill = navy_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    print(f"Excel workbook successfully generated: {output_filename}")

if __name__ == "__main__":
    generate_styled_excel()
