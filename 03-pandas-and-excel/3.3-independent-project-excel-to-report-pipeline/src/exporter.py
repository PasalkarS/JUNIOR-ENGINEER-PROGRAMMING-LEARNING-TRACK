"""
OpenPyXL-powered multi-worksheet professional Excel generator.
"""
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_executive_report(
    enriched_df: pd.DataFrame,
    category_summary: pd.DataFrame,
    regional_summary: pd.DataFrame,
    audit_errors: pd.DataFrame,
    output_path: str | Path
) -> Path:
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        category_summary.to_excel(writer, sheet_name="Category Summary", index=False)
        regional_summary.to_excel(writer, sheet_name="Regional Summary", index=False)
        enriched_df.to_excel(writer, sheet_name="Clean Sales Data", index=False)
        if not audit_errors.empty:
            audit_errors.to_excel(writer, sheet_name="Audit Rejections", index=False)

        wb = writer.book

        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            for col in ws.columns:
                max_len = max(len(str(c.value or "")) for c in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    return out
